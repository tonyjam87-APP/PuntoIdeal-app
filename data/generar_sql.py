# Lee el Excel (valores calculados) y genera supabase_congestion.sql
import openpyxl, math, os
d=os.path.dirname(__file__)
wb=openpyxl.load_workbook(os.path.join(d,"PuntoIdeal_modelo_v10.xlsx"),data_only=True)
q=lambda v:"null" if v is None or v=="" else (("'"+str(v).replace("'","''")+"'") if isinstance(v,str) else str(v))
def mover(lat,lon,rumbo,m):
    b=math.radians(rumbo); return round(lat+m*math.cos(b)/111320,6), round(lon+m*math.sin(b)/(111320*math.cos(math.radians(lat))),6)
out=["""-- PoC congestión en Av. Perón (semáforos, paradas y desvíos)
drop table if exists congestion; drop table if exists paradas; drop table if exists desvios;
create table congestion (
  id int primary key, nombre text not null, categoria text not null,
  latitud double precision not null, longitud double precision not null, valor numeric not null,
  semaforo text, sentido text, hora int,
  autos_hora numeric, paradas_afectan numeric, verde_perdido numeric, capacidad numeric,
  cola numeric, pct_desvio numeric, autos_hora_desvio numeric,
  indice_desvio numeric, cola_desvio numeric, mejora numeric, categoria_desvio text
);
create table paradas (id int primary key, nombre text, lineas text, latitud double precision, longitud double precision, nota text);
create table desvios (id int primary key, nombre text, quien text, recorrido text, sentido text, alivia_197 text, alivia_pueyrredon text,
  porcentaje numeric, origen_lat double precision, origen_lon double precision, destino_lat double precision, destino_lon double precision, nota text);
alter table congestion enable row level security; alter table paradas enable row level security; alter table desvios enable row level security;
create policy "lectura publica" on congestion for select using (true);
create policy "lectura publica" on paradas for select using (true);
create policy "lectura publica" on desvios for select using (true);
"""]
C=wb["Congestión"]
for r in range(2,C.max_row+1):
    g=lambda c:C.cell(r,c).value
    if g(1) is None: continue
    vals=[g(1),g(2),g(24),g(6),g(7),g(17),g(3),g(4),g(5),g(8),g(10),g(15),g(16),g(18),g(19),g(20),g(21),g(22),g(23),None]
    ind=g(21); cd="Saturado" if ind>=100 else "Cargado" if ind>=80 else "Fluido"
    vals[-1]=cd
    vals[2]=g(24)
    out.append("insert into congestion values (%s);"%",".join(q(v) for v in vals))
Pa=wb["Paradas de colectivo"]
for r in range(2,10):
    out.append("insert into paradas values (%s);"%",".join(q(v) for v in (Pa.cell(r,1).value,Pa.cell(r,2).value,Pa.cell(r,3).value,Pa.cell(r,5).value,Pa.cell(r,6).value,Pa.cell(r,7).value)))
c197=(-34.516764,-58.746806); pue=(-34.520418,-58.741839); lav=(-34.518621,-58.744304)
cb=(round((-34.522520-34.522416)/2,6),round((-58.739016-58.739171)/2,6))
D=wb["Desvíos"]
geo={1:(mover(*c197,235,90),cb),2:(pue,mover(*pue,222,150)),3:(lav,mover(*lav,222,100)),4:(pue,mover(*pue,42,130))}
nota="Recorrido esquemático: los puntos son aproximados, no siguen la calle exacta."
for i in range(1,5):
    r=i+1; (olat,olon),(dlat,dlon)=geo[i]
    out.append("insert into desvios values (%s);"%",".join(q(v) for v in (i,D.cell(r,2).value,D.cell(r,3).value,D.cell(r,4).value,D.cell(r,6).value,D.cell(r,7).value,D.cell(r,8).value,D.cell(r,9).value,olat,olon,dlat,dlon,nota)))
open(os.path.join(d,"supabase_congestion.sql"),"w",encoding="utf8").write("\n".join(out)+"\n")
print(len(out))
